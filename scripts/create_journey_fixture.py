"""Synthetic Canvas course for a disposable-account journey; no student materials."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
from docx import Document
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parents[1] / 'output/journey'
root.mkdir(parents=True, exist_ok=True)
lessons = [
    ('STU5101 Week 1 - Research Questions', [
        ('Lecture objectives', 'Distinguish a hypothesis from a research question. A research question asks what relationship or process should be investigated. A hypothesis is a specific, testable prediction about the relationship between variables. A theory explains a wider set of observations. A hypothesis may be derived from a theory but is not the theory itself.'),
        ('Variables and operational definitions', 'A variable is an attribute that can take different values. An independent variable is manipulated or used to predict a change. A dependent variable is the outcome measured. An operational definition specifies exactly how an abstract concept is measured. For example, defining engagement as minutes spent on a learning task makes the measurement explicit.'),
        ('Classroom activity', 'A researcher predicts that students who practice retrieval will remember more concepts than students who only reread notes. Study method is the independent variable, and number of concepts recalled is the dependent variable. Correlation alone cannot establish causation. A confounding variable is an alternative factor related to both the predictor and the outcome. Prior knowledge could confound an observational comparison between study groups.'),
    ]),
    ('STU5101 Week 2 - Research Design', [
        ('Lecture objectives', 'Compare experimental and observational research. An experiment manipulates an independent variable and measures a dependent variable. An observational study records existing conditions without manipulating them. Random assignment allocates participants to conditions by chance. Random sampling selects people from a population. These two procedures solve different problems.'),
        ('Validity and reliability', 'Internal validity concerns whether a causal explanation is credible within a study. External validity concerns whether findings generalise to other people or settings. Reliability is consistency of measurement, while validity is whether the measurement captures the intended concept. A bathroom scale that is always five kilograms too high can be reliable but inaccurate.'),
        ('Tutorial example', 'To test a new study method, randomly assign students to retrieval practice or rereading. Give both groups the same material and equal study time, and use the same final test. Random assignment helps reduce pre-existing group differences. A control group provides a comparison. Blinding helps reduce expectation bias. A sample from one small class may limit external validity even if the experiment has strong internal validity.'),
    ]),
    ('STU5101 Week 3 - Effective Learning', [
        ('Lecture objectives', 'Use retrieval practice, spaced practice and interleaving to improve learning. Retrieval practice means attempting to recall information from memory rather than simply reading it again. Rereading can create familiarity without demonstrating recall. A student should first attempt a question without looking at the answer and then check feedback.'),
        ('Spacing and interleaving', 'Spaced practice distributes learning across separate sessions instead of concentrating it into one long session. Interleaving mixes different but related problem types rather than completing one type in a single block. Spacing concerns when practice happens; interleaving concerns which problems are mixed. These strategies can be combined in the same study plan.'),
        ('Feedback and transfer', 'Corrective feedback identifies the correct answer and explains why a response was wrong. Transfer means applying knowledge to a new situation. A transfer question asks the student to use a principle in an unfamiliar example. Metacognition is monitoring and regulating one\'s own learning. Confidence should be checked against actual performance, not only a feeling that notes look familiar.'),
    ]),
]

def pdf(path, title, sections):
    c = canvas.Canvas(str(path))
    c.setTitle(title)
    for heading, text in sections:
        c.setFont('Helvetica-Bold', 16)
        c.drawString(45, 790, title)
        c.setFont('Helvetica-Bold', 14)
        c.drawString(45, 752, heading)
        c.setFont('Helvetica', 12)
        for i, line in enumerate(simpleSplit(text, 'Helvetica', 12, 500)):
            c.drawString(45, 716 - i * 20, line)
        c.showPage()
    c.save()

p1 = root / 'Week 1 Lecture.pdf'
pdf(p1, *lessons[0])
prs = Presentation()
for heading, text in lessons[1][1]:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = lessons[1][0] + '\n' + heading
    slide.placeholders[1].text = text
p2 = root / 'Week 2 Lecture.pptx'
prs.save(str(p2))
doc = Document()
doc.add_heading(lessons[2][0], 0)
for heading, text in lessons[2][1]:
    doc.add_heading(heading, 1)
    doc.add_paragraph(text)
p3 = root / 'Week 3 Lecture Notes.docx'
doc.save(str(p3))
reading = root / 'Optional Journal Reading.pdf'
pdf(reading, 'Journal of Unrelated Studies - Supplementary Reading', [('Research article - not teaching slides', 'Abstract. This optional research article discusses historical astronomy. Authors: Example Author. Journal volume 12 issue 4. DOI 10.0000/example. Methods, results, discussion, and references are research article sections. Its focus is intentionally unrelated to STU5101 classroom concepts. Exclude this extracurricular reading from classroom-based teaching and practice.')])
archive = root / 'STU5101 - Whole Course Canvas.zip'
with ZipFile(archive, 'w', ZIP_DEFLATED) as z:
    for path, name in [(p1, 'STU5101/Week 1/lecture.pdf'), (p2, 'STU5101/Week 2/slides.pptx'), (p3, 'STU5101/Week 3/lecture-notes.docx'), (p1, 'STU5101/Duplicates/lecture-copy.pdf'), (reading, 'STU5101/Readings/optional-journal.pdf')]:
        z.write(path, name)
    z.writestr('STU5101/Canvas-links.html', '<p>Canvas export links; not a lecture.</p>')
    z.writestr('__MACOSX/._lecture.pdf', 'Ignored system metadata')
print(archive)
